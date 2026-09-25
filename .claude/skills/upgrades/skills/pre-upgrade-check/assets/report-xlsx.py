#!/usr/bin/env python3
"""Render a pre-upgrade-check report.csv into a styled report.xlsx.

A CSV carries no formatting, so importing it never reproduces the team sheet's look. This makes a
workbook for import (File -> Import -> Upload) with a dark header, bold white centered
titles, thin borders, a frozen header row, wrapped cells, and sized columns.

Usage:  python3 report-xlsx.py <report.csv> [report.xlsx]
        (default output: report.xlsx next to the input)

Requires openpyxl:  python3 -m pip install --user openpyxl
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Column widths keyed by the header text (falls back to a default for anything else).
WIDTHS = {
    "#": 5,
    "Description": 48,
    "Page": 24,
    "Evidence": 42,
    "Analysis": 62,
    "Response": 20,
}


def build(src: Path, out: Path) -> None:
    rows = list(csv.reader(src.open(newline="")))
    if not rows:
        raise SystemExit(f"{src}: empty")

    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"

    header_fill = PatternFill("solid", fgColor="434343")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    data_font = Font(size=10)
    thin = Side(style="thin", color="C0C0C0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    num_align = Alignment(horizontal="center", vertical="top")

    for r in rows:
        ws.append(r)

    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = header_fill, header_font, header_align, border
    ws.row_dimensions[1].height = 30

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i, c in enumerate(row):
            c.font, c.border = data_font, border
            c.alignment = num_align if i == 0 else cell_align

    from openpyxl.utils import get_column_letter
    for i, head in enumerate(rows[0], start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(head, 24)

    ws.freeze_panes = "A2"
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".xlsx")
    build(src, out)
